"""
Extraction automatique d'une nomenclature de pièces à partir d'une présentation de conception (.pptx).

Deux modes :
  --sans-ia : lit uniquement les tableaux de la présentation (rapide, gratuit, mais ignore le texte libre).
  par défaut : envoie le contenu des diapositives à un modèle de langage (Claude) qui applique aussi
               les évolutions écrites en texte libre (ajouts, remplacements, changements de quantité).

Dans les deux cas, le résultat passe par des contrôles automatiques et chaque ligne douteuse est
marquée « AVerifier = oui » : l'IA propose, l'équipe métier valide (dans l'application Power Apps).

Exemples :
  python extraction_nomenclature.py ../donnees/presentation_conception_exemple.pptx --sans-ia
  python extraction_nomenclature.py ../donnees/presentation_conception_exemple.pptx \
         --comparer ../donnees/nomenclature_attendue.csv
"""

import argparse
import csv
import json
import os
import re
import sys

from pptx import Presentation

COLONNES = ["Reference", "Designation", "Quantite", "SousEnsemble", "Poste", "Statut", "AVerifier", "Commentaire"]
MOTIF_REFERENCE = re.compile(r"^[A-Z0-9]+(-[A-Z0-9]+)+$")
MODELE_PAR_DEFAUT = os.environ.get("MODELE_IA", "claude-sonnet-5-5")


# --------------------------------------------------------------------------- lecture du PowerPoint
def lire_presentation(chemin):
    """Retourne une liste de diapositives : {"numero", "textes": [...], "tableaux": [[ligne, ...], ...]}."""
    diapositives = []
    for numero, diapo in enumerate(Presentation(chemin).slides, start=1):
        textes, tableaux = [], []
        for forme in diapo.shapes:
            if forme.has_table:
                tableaux.append([[cellule.text.strip() for cellule in ligne.cells] for ligne in forme.table.rows])
            elif forme.has_text_frame:
                for paragraphe in forme.text_frame.paragraphs:
                    texte = "".join(run.text for run in paragraphe.runs).strip()
                    if texte:
                        textes.append(texte)
        diapositives.append({"numero": numero, "textes": textes, "tableaux": tableaux})
    return diapositives


def en_texte(diapositives):
    """Met en forme le contenu des diapositives pour le modèle de langage."""
    blocs = []
    for d in diapositives:
        lignes = [f"--- Diapositive {d['numero']} ---"] + d["textes"]
        for tableau in d["tableaux"]:
            lignes += [" | ".join(ligne) for ligne in tableau]
        blocs.append("\n".join(lignes))
    return "\n\n".join(blocs)


# --------------------------------------------------------------------------- mode sans IA
def normaliser(entete):
    entete = entete.lower()
    for a, b in (("é", "e"), ("è", "e"), ("ê", "e"), ("à", "a"), ("ô", "o"), (" ", ""), ("-", "")):
        entete = entete.replace(a, b)
    return entete


CORRESPONDANCES = {
    "reference": "Reference", "ref": "Reference",
    "designation": "Designation", "libelle": "Designation",
    "qte": "Quantite", "quantite": "Quantite",
    "sousensemble": "SousEnsemble",
    "poste": "Poste", "postedemontage": "Poste",
}


def extraire_tableaux(diapositives):
    pieces = []
    for d in diapositives:
        for tableau in d["tableaux"]:
            if not tableau:
                continue
            colonnes = [CORRESPONDANCES.get(normaliser(c)) for c in tableau[0]]
            if "Reference" not in colonnes:
                continue
            for ligne in tableau[1:]:
                piece = {c: v for c, v in zip(colonnes, ligne) if c}
                piece["Statut"] = "issue du tableau"
                piece["Commentaire"] = f"Diapositive {d['numero']}"
                pieces.append(piece)
    return pieces


# --------------------------------------------------------------------------- mode IA
CONSIGNE_SYSTEME = """Tu es un assistant d'ingénierie industrielle. On te donne le contenu d'une présentation
de revue de conception. Construis la nomenclature FINALE des pièces :
- pars des tableaux de nomenclature ;
- applique toutes les évolutions décrites en texte libre (ajouts, remplacements, changements de quantité) ;
- une pièce remplacée ne doit plus apparaître ; sa remplaçante porte le statut "remplace <ancienne référence>".
Réponds UNIQUEMENT avec un objet JSON, sans texte autour, au format :
{"pieces": [{"Reference": "...", "Designation": "...", "Quantite": 1, "SousEnsemble": "...", "Poste": "...",
  "Statut": "inchangée | ajoutée | quantité modifiée | remplace <ref>", "Commentaire": "diapositive source"}],
 "points_a_verifier": ["informations manquantes ou ambiguës"]}
N'invente jamais une information absente : laisse le champ vide et signale-le dans points_a_verifier."""


def extraire_avec_ia(diapositives, modele):
    try:
        import anthropic
    except ImportError:
        sys.exit("Installez le SDK : pip install anthropic (ou utilisez --sans-ia).")
    if not os.environ.get("ANTHROPIC_API_KEY"):
        sys.exit("Définissez la variable d'environnement ANTHROPIC_API_KEY (ou utilisez --sans-ia).")

    client = anthropic.Anthropic()
    reponse = client.messages.create(
        model=modele,
        max_tokens=4000,
        system=CONSIGNE_SYSTEME,
        messages=[{"role": "user", "content": en_texte(diapositives)}],
    )
    texte = "".join(bloc.text for bloc in reponse.content if bloc.type == "text")
    texte = re.sub(r"^```(json)?|```$", "", texte.strip(), flags=re.MULTILINE).strip()
    donnees = json.loads(texte)
    for point in donnees.get("points_a_verifier", []):
        print(f"  ⚠ Point signalé par l'IA : {point}")
    return donnees.get("pieces", [])


# --------------------------------------------------------------------------- contrôles
def controler(pieces):
    """Ajoute AVerifier = oui et une explication pour chaque ligne qui ne passe pas les contrôles."""
    vues = {}
    for piece in pieces:
        problemes = []
        ref = str(piece.get("Reference", "")).strip()
        if not MOTIF_REFERENCE.match(ref):
            problemes.append("référence absente ou mal formée")
        try:
            brut = piece.get("Quantite")
            quantite = int(str("" if brut is None else brut).strip())
            if quantite <= 0:
                problemes.append("quantité nulle ou négative")
            piece["Quantite"] = quantite
        except ValueError:
            problemes.append("quantité non numérique")
        for champ in ("Designation", "SousEnsemble", "Poste"):
            if not str(piece.get(champ, "")).strip():
                problemes.append(f"{champ} manquant")
        if ref in vues:
            problemes.append("référence en double")
        vues[ref] = True
        piece["AVerifier"] = "oui" if problemes else "non"
        if problemes:
            piece["Commentaire"] = (str(piece.get("Commentaire", "")) + " | " + ", ".join(problemes)).strip(" |")
    return pieces


def comparer(pieces, chemin_attendu):
    with open(chemin_attendu, encoding="utf-8-sig") as f:
        attendu = {l["Reference"]: l for l in csv.DictReader(f)}
    obtenu = {str(p.get("Reference")): p for p in pieces}
    justes = [r for r in attendu if r in obtenu and str(obtenu[r].get("Quantite")) == attendu[r]["Quantite"]]
    print(f"\nComparaison : {len(justes)}/{len(attendu)} pièces retrouvées avec la bonne quantité.")
    for r in attendu:
        if r not in justes:
            print(f"  ✗ {r} : attendu {attendu[r]['Quantite']}, obtenu {obtenu.get(r, {}).get('Quantite') if obtenu.get(r, {}).get('Quantite') not in (None, '') else ('vide' if r in obtenu else 'absente')}")
    for r in obtenu:
        if r not in attendu:
            print(f"  ✗ {r} : pièce en trop")


# --------------------------------------------------------------------------- programme principal
def main():
    parser = argparse.ArgumentParser(description="Extraction de nomenclature depuis un PowerPoint de conception.")
    parser.add_argument("presentation", help="chemin du fichier .pptx")
    parser.add_argument("--sortie", default="nomenclature.csv", help="fichier CSV produit (défaut : nomenclature.csv)")
    parser.add_argument("--sans-ia", action="store_true", help="lire uniquement les tableaux, sans modèle d'IA")
    parser.add_argument("--modele", default=MODELE_PAR_DEFAUT, help="modèle Claude à utiliser")
    parser.add_argument("--comparer", help="CSV de nomenclature attendue, pour mesurer la qualité du résultat")
    args = parser.parse_args()

    diapositives = lire_presentation(args.presentation)
    print(f"{len(diapositives)} diapositives lues dans {args.presentation}")

    if args.sans_ia:
        pieces = extraire_tableaux(diapositives)
        print("Mode sans IA : seules les pièces des tableaux sont reprises (les évolutions en texte libre sont ignorées).")
    else:
        print(f"Mode IA ({args.modele}) : tableaux + évolutions en texte libre.")
        pieces = extraire_avec_ia(diapositives, args.modele)

    pieces = controler(pieces)
    with open(args.sortie, "w", newline="", encoding="utf-8-sig") as f:  # utf-8-sig : accents corrects dans Excel
        ecrivain = csv.DictWriter(f, fieldnames=COLONNES, extrasaction="ignore")
        ecrivain.writeheader()
        ecrivain.writerows(pieces)

    a_verifier = sum(p["AVerifier"] == "oui" for p in pieces)
    print(f"{len(pieces)} pièces écrites dans {args.sortie} — {a_verifier} à vérifier par l'équipe métier.")
    if args.comparer:
        comparer(pieces, args.comparer)


if __name__ == "__main__":
    main()
